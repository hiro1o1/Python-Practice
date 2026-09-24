"""Add the 10.1" Self-service Price Checker page to the catalog PPTX.

    python catalog/tools/add_price_checker.py <in.pptx> <out.pptx>

The new page goes after "Bar Type and OLED Displays" (Digital Signage). Every
shape is cloned from an existing product page so fonts, colours and spacing
match. Existing pages are only touched where the insert shifts numbers:
slide-number fields, the page references on the portfolio page, and the
"In this section" lists on the section openers (opener 02 also lists the
new page).

Specs come from the supplier sheet "10.1Machine.doc" (model XYC-CJ-D5).
"""

import copy
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PHOTOS = os.path.join(os.path.dirname(HERE), "photos")
ICONS = os.path.join(HERE, "icons")

AFTER_TITLE = "Bar Type and OLED Displays"
TITLE = '10.1" Self-service Price Checker'
SUBTITLE = "Barcode price checker for retail  ·  Model XYC-CJ-D5"
PICTURES = [  # (file, alt text, caption)
    ("price-checker-front.png", "10.1 inch self-service price checker, front view", "Front view"),
    ("price-checker-scanner.png", "10.1 inch self-service price checker, scanner window", "Scanner window"),
]
CARDS = [  # (value, note, icon)
    ('10.1"', "IPS, 1280 × 800", "monitor"),
    ("1D / 2D", "Barcode scanning", "barcode"),
    ("Android", "Android 10", "cpu"),
    ("IP54", "Dust and splash", "shield"),
]
FEATURES = [
    ["Scan any barcode to see price and promotion",
     "Live prices and stock from ERP / POS",
     "Multi-language; multi-currency on some models",
     "Touch and scan; voice price announcement"],
    ["Promotions, member prices, QR member login",
     "Product details: ingredients, usage, origin",
     "Top-searched products and peak times",
     "Remote fault reporting and maintenance"],
]
SPECS = [
    ("Display", "10.1\" IPS LCD, 1280 × 800, 250–300 cd/m², 60 Hz"),
    ("Touch", "10-point touch, 2 mm tempered glass"),
    ("Platform", "A133 quad-core ARM Cortex-A53, 1.5 GHz · Android 10"),
    ("Memory", "2 GB RAM, 32 GB storage"),
    ("Scanner", "1D and 2D codes: EAN-13, UPC, Code 128, QR Code"),
    ("Connectivity", "Wi-Fi 2.4 GHz, Bluetooth 5.0, RJ45 LAN, USB 2.0 + 3.0"),
    ("Power", "DC 12 V adapter, max. 15 W, standby < 0.5 W"),
    ("Size and weight", "267 × 272 × 89 mm, 1.2 kg · wall bracket included"),
    ("Environment", "Operating 0–40 °C, storage −20–60 °C"),
]
OPTIONS = "Options on request: camera (720P or 1080P), NFC / RFID, PoE power, other chip and memory configurations."
USES = ["Supermarkets", "Shopping malls", "Convenience stores", "Warehouses"]

SEMIBOLD = os.path.expanduser("~/.fonts/SegoeUI-Semibold-subset.ttf")


# --------------------------------------------------------------------------- helpers
def find(prs, title=None, name=None):
    for s in prs.slides:
        for sh in s.shapes:
            if title and sh.name == "Page title" and sh.text_frame.text.replace("\n", " ") == title:
                return s
            if name and sh.name == name:
                return sh
    raise KeyError(title or name)


def shape_on(slide, name, text=None):
    for sh in slide.shapes:
        if sh.name == name and (text is None or (sh.has_text_frame and sh.text_frame.text.startswith(text))):
            return sh
    raise KeyError(name)


class Page:
    def __init__(self, slide):
        self.slide = slide
        self.tree = slide.shapes._spTree
        self.next_id = 100

    def clone(self, src, x=None, y=None, w=None, h=None, text=None, rename=None):
        el = copy.deepcopy(src._element)
        self.next_id += 1
        nv = el.find(".//" + qn("p:cNvPr"))
        nv.set("id", str(self.next_id))
        if rename:
            nv.set("name", rename)
        xfrm = el.find(".//" + qn("a:xfrm"))
        if xfrm is None:
            xfrm = el.find(".//" + qn("p:xfrm"))
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        if x is not None:
            off.set("x", str(int(Mm(x))))
            off.set("y", str(int(Mm(y))))
        if w is not None:
            ext.set("cx", str(int(Mm(w))))
            ext.set("cy", str(int(Mm(h))))
        if text is not None:
            set_text(el, text)
        self.tree.append(el)
        return el


def set_text(el, text):
    """Replace a text body with one paragraph per line, keeping the first paragraph's formatting."""
    lines = text if isinstance(text, list) else [text]
    paras = el.findall(".//" + qn("a:p"))
    body = paras[0].getparent()
    tpl = paras[0]
    for p in paras:
        body.remove(p)
    for line in lines:
        p = copy.deepcopy(tpl)
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = line
        body.append(p)


def text_width_mm(text, pt):
    f = ImageFont.truetype(SEMIBOLD, 1000)
    return f.getlength(text) / 1000 * pt * 25.4 / 72


# --------------------------------------------------------------------------- page
def build(prs):
    bar = find(prs, title=AFTER_TITLE)
    ops = find(prs, title="OPS PC Module")
    smart = find(prs, title="4K AI Smartboards")
    slide = prs.slides.add_slide(bar.slide_layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bar.background.fill.fore_color.rgb
    pg = Page(slide)

    for n in ("Running header", "Section label", "Footer", "Page number"):
        pg.clone(shape_on(bar, n))
    pg.clone(shape_on(bar, "Page title"), text=TITLE)
    pg.clone(shape_on(bar, "Subtitle"), text=SUBTITLE)

    # Pictures in two white panels with captions.
    panel = shape_on(ops, "Image panel")
    caption = shape_on(ops, "Caption")
    for i, (fname, alt, cap) in enumerate(PICTURES):
        x = 16 + i * 92
        pg.clone(panel, x, 46, 86, 56)
        path = os.path.join(PHOTOS, fname)
        from PIL import Image
        iw, ih = Image.open(path).size
        bw, bh = 76, 48
        s = min(bw / iw, bh / ih)
        w, h = iw * s, ih * s
        pic = slide.shapes.add_picture(path, Mm(x + 5 + (bw - w) / 2), Mm(46 + 4 + (bh - h) / 2), Mm(w), Mm(h))
        pic.name = alt
        pic._element.find(".//" + qn("p:cNvPr")).set("descr", alt)
        pg.clone(caption, x, 104, 86, 4, text=cap)

    # Stat cards with icons (card geometry as on the smartboard page).
    card, value, note = (shape_on(smart, n) for n in ("Card", "Card value", "Card note"))
    icon_src = next(sh for sh in smart.shapes if sh.name.startswith("Icon"))
    for i, (v, nt, icon) in enumerate(CARDS):
        x = 16 + i * 45.75
        pg.clone(card, x, 111, 40.8, 21)
        pg.clone(value, x + 5, 114.5, 30.8, 8, text=v)
        pg.clone(note, x + 5, 124, 30.8, 5, text=nt)
        d = icon_src.width
        ic = slide.shapes.add_picture(os.path.join(ICONS, icon + ".png"),
                                      Mm(x + 40.8 - 3.2) - d, Mm(114.5 + 4) - d // 2, d, d)
        ic.name = f"Icon {icon}"

    # Key features (two columns), cloned from the smartboard page.
    label = shape_on(smart, "Label", "Key features")
    feat = [sh for sh in smart.shapes if sh.name == "Key features"]
    pg.clone(label, 16, 137, 178, 4, text="Key features")
    for i, col in enumerate(FEATURES):
        pg.clone(feat[i], 16 + i * 94, 143, 84, 25, text=col)

    # Specifications table, cloned from the OPS page (same row count).
    table_src = next(sh for sh in ops.shapes if sh.has_table)
    el = pg.clone(table_src, 16, 172, 178, 67.5, rename="Price checker specifications")
    tbl = el.find(".//" + qn("a:tbl"))
    rows = tbl.findall(qn("a:tr"))
    assert len(rows) == len(SPECS), "spec rows must match the cloned table"
    for tr, (k, v) in zip(rows, SPECS):
        tr.set("h", str(int(Mm(7.5))))
        cells = tr.findall(qn("a:tc"))
        set_text(cells[0].find(qn("a:txBody")), k)
        set_text(cells[1].find(qn("a:txBody")), v)
    ext = el.find(".//" + qn("p:xfrm")).find(qn("a:ext"))
    ext.set("cy", str(int(Mm(7.5)) * len(SPECS)))
    pg.clone(shape_on(smart, "Card note"), 16, 243, 178, 5, text=OPTIONS, rename="Options note")

    # Applications chips, as on the Bar Type page.
    chip = next(sh for sh in bar.shapes if sh.name.startswith("Chip"))
    pg.clone(shape_on(bar, "Label", "Applications"), 16, 253, 100, 4, text="Applications")
    size = Pt(11)
    for sh in bar.shapes:
        if sh.name.startswith("Chip"):
            r = sh._element.find(".//" + qn("a:rPr"))
            size = int(r.get("sz", "1100")) / 100
            break
    x = 16
    for u in USES:
        w = text_width_mm(u, size) + 9
        pg.clone(chip, x, 260, w, 9.4, text=u, rename=f"Chip {u}")
        x += w + 2.8
    return slide


def renumber(prs, new_slide):
    order = list(prs.slides)
    pos = order.index(new_slide) + 1  # 1-based page of the new slide
    for i, s in enumerate(order, 1):
        for f in s.shapes._spTree.iter(qn("a:fld")):
            if f.get("type") == "slidenum":
                f.find(qn("a:t")).text = str(i)
        for sh in s.shapes:
            if sh.name == "Page reference" and sh.has_text_frame:
                run = sh.text_frame.paragraphs[0].runs[0]
                a, b = run.text.replace("Page ", "").split("–")
                a, b = int(a), int(b)
                if b >= pos - 1 and a <= pos - 1:  # the section that gains the page
                    b += 1
                elif a >= pos:
                    a, b = a + 1, b + 1
                run.text = f"Page {a:02d}–{b:02d}"
            elif sh.name.startswith("Section page ") and sh.has_text_frame:
                run = sh.text_frame.paragraphs[0].runs[0]
                n = int(run.text)
                if n >= pos:
                    run.text = f"{n + 1:02d}"
    return pos


def add_to_opener(prs, pos):
    """List the new page under "In this section" on its section's opener."""
    order = list(prs.slides)
    for s in reversed(order[:pos - 1]):
        nums = [sh for sh in s.shapes if sh.name.startswith("Section page ")]
        if not nums:
            continue
        last = max(nums, key=lambda sh: int(sh.text_frame.text))
        # Title box: the text box on the same line, to the right of the number.
        title = min((sh for sh in s.shapes if sh.has_text_frame and sh is not last
                     and abs(sh.top - last.top) < Mm(1) and sh.left > last.left),
                    key=lambda sh: sh.left - last.left)
        col_x = sorted({sh.left for sh in nums})
        rows = sorted({sh.top for sh in nums})
        step = rows[1] - rows[0] if len(rows) > 1 else Mm(7.5)
        in_last_col = [sh for sh in nums if sh.left == last.left]
        y = last.top + step
        pg = Page(s)
        pg.next_id = 900
        n_el = pg.clone(last, rename=f"Section page {pos}")
        t_el = pg.clone(title, rename="Section page title")
        for e in (n_el, t_el):
            off = e.find(".//" + qn("a:off"))
            off.set("y", str(int(y)))
        set_text(n_el.find(qn("p:txBody")), f"{pos:02d}")
        set_text(t_el.find(qn("p:txBody")), TITLE)
        return s
    return None


def main(src, dst):
    prs = Presentation(src)
    bar = find(prs, title=AFTER_TITLE)
    slide = build(prs)
    lst = prs.slides._sldIdLst
    rid = next(r.rId for r in prs.part.rels.values() if r._target is slide.part)
    new_el = next(e for e in lst if e.get(qn("r:id")) == rid)
    bar_rid = next(r.rId for r in prs.part.rels.values() if r._target is bar.part)
    bar_el = next(e for e in lst if e.get(qn("r:id")) == bar_rid)
    lst.remove(new_el)
    bar_el.addnext(new_el)
    pos = renumber(prs, slide)
    add_to_opener(prs, pos)
    prs.save(dst)
    print("saved", dst, "new page", pos, "of", len(prs.slides))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
