"""Make emails and phone numbers clickable (mailto: and WhatsApp) in the catalog PPTX.

    python catalog/tools/add_contact_links.py <in.pptx> <out.pptx>
"""

import copy
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE = re.compile(r"\+86 \d{3} \d{4} \d{4}")


def wa(num):
    return "https://wa.me/" + re.sub(r"\D", "", num)


def link_run(slide, r, url):
    rpr = r._r.get_or_add_rPr()
    for old in rpr.findall(qn("a:hlinkClick")):
        rpr.remove(old)
    rid = slide.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
    h = rpr.makeelement(qn("a:hlinkClick"), {qn("r:id"): rid})
    # hlinkClick must follow fill/font children in rPr
    rpr.append(h)


def split_run(r, spans):
    """Split run r into pieces [(text, url|None)], keeping its formatting."""
    parent = r._r.getparent()
    idx = list(parent).index(r._r)
    out = []
    for text, url in spans:
        el = copy.deepcopy(r._r)
        el.find(qn("a:t")).text = text
        parent.insert(idx, el)
        idx += 1
        out.append((el, url))
    parent.remove(r._r)
    return out


def main(src, dst):
    prs = Presentation(src)
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in list(para.runs):
                    t = r.text
                    hits = [(m.start(), m.end(), "mailto:" + m.group()) for m in EMAIL.finditer(t)] + \
                           [(m.start(), m.end(), wa(m.group())) for m in PHONE.finditer(t)]
                    if not hits:
                        continue
                    rpr = r._r.find(qn("a:rPr"))
                    if rpr is not None and rpr.find(qn("a:hlinkClick")) is not None and len(hits) == 1 \
                            and hits[0][0] == 0 and hits[0][1] == len(t):
                        continue  # already linked
                    hits.sort()
                    pieces, pos = [], 0
                    for a, b, url in hits:
                        if a > pos:
                            pieces.append((t[pos:a], None))
                        pieces.append((t[a:b], url))
                        pos = b
                    if pos < len(t):
                        pieces.append((t[pos:], None))
                    for el, url in split_run(r, pieces):
                        if url:
                            from pptx.text.text import _Run
                            link_run(s, _Run(el, para), url)
                            n += 1
    prs.save(dst)
    print("links added:", n)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
