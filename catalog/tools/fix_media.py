"""Fix specific images inside the catalog PPTX (media files only; slides untouched).

    python catalog/tools/fix_media.py <in.pptx> <out.pptx>

* white-box product photos -> transparent background (no box edge in white cards)
* hazy "floor reflection" shadows baked into cut-outs -> removed (they show as
  clouds in some PDF viewers); the slide's own soft shadow stays
"""

import io
import sys
import zipfile

import cv2
import numpy as np
from PIL import Image
from pptx import Presentation

sys.path.insert(0, __import__("os").path.dirname(__file__))
from fix_edges import decontaminate  # noqa: E402

WHITE_BOX = {22: None}          # every white-background picture on these pages
HAZE = {9: ("Classroom board",)}  # pictures whose names start with these


def knock_out_white(rgb, thr=248):
    x = rgb.astype(np.int32)
    w, sat = x.min(-1), x.max(-1) - x.min(-1)
    cand = ((w >= thr) & (sat < 10)).astype(np.uint8)
    n, lab = cv2.connectedComponents(cand)
    border = set(np.unique(np.r_[lab[0], lab[-1], lab[:, 0], lab[:, -1]])) - {0}
    bg = np.isin(lab, list(border))
    a = np.where(bg, 0, 255).astype(np.uint8)
    a = cv2.GaussianBlur(a, (3, 3), 0)
    a[~bg] = 255
    return decontaminate(np.dstack([rgb, a]))


def remove_haze(rgba, keep=200):
    a = rgba[..., 3]
    solid = (a >= keep).astype(np.uint8)
    n, lab, st, _ = cv2.connectedComponentsWithStats(solid)
    solid = (lab == 1 + np.argmax(st[1:, 4])).astype(np.uint8)
    ff = solid.copy()
    h, w = solid.shape
    cv2.floodFill(ff, np.zeros((h + 2, w + 2), np.uint8), (0, 0), 1)
    solid = solid | (1 - ff)
    na = cv2.GaussianBlur(solid.astype(np.float32) * 255, (3, 3), 0).astype(np.uint8)
    out = rgba.copy()
    out[..., 3] = na
    return decontaminate(out)


def main(src, dst):
    prs = Presentation(src)
    new = {}
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.shape_type != 13:
                continue
            part = s.part.related_part(sh._element.blipFill.blip.rEmbed)
            name = str(part.partname).lstrip("/")
            im = Image.open(io.BytesIO(part.blob))
            arr = np.asarray(im.convert("RGBA"))
            if i in WHITE_BOX and arr[2, 2, :3].min() > 240 and arr[2, 2, 3] == 255:
                new[name] = knock_out_white(arr[..., :3])
            elif i in HAZE and sh.name.startswith(HAZE[i]):
                new[name] = remove_haze(arr)
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in new:
                buf = io.BytesIO()
                Image.fromarray(new[item.filename], "RGBA").save(buf, "PNG", optimize=True)
                data = buf.getvalue()
                print("fixed", item.filename)
            zout.writestr(item, data)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
